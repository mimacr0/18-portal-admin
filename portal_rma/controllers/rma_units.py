# -*- coding: utf-8 -*-
import json
import io
import math
import logging
import pytz
from datetime import datetime

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import xlsxwriter
except ImportError:
    xlsxwriter = None

from odoo import http, _, fields
from odoo.http import request
from odoo.osv import expression
from odoo.addons.portal_account.controllers.dashboard import PortalDashboardController

_logger = logging.getLogger(__name__)

class PortalRmaController(PortalDashboardController):

    RMA_UNIT_FIELDS_MAPPING = {
        'name': 'name',
        'product': 'product_id',
        'serial': 'serial',
        'imei': 'imei',
        'location': 'location_id',
        'state': 'state',
        'condition': 'condition',
        'create_date': 'create_date',
        'account_sku': 'account_sku',
        'account_ean13': 'account_ean13',
        'account_fnsku': 'account_fnsku',
        'account_asin': 'account_asin',
    }

    def _get_admin_layout_menus(self):
        menus = super()._get_admin_layout_menus()
        menus.append({
            'id': 'rma_units',
            'name': _('RMA Units'),
            'url': '/account/rma/units',
            'icon': 'fas fa-boxes',
            'order': 30
        })
        return menus

    def _get_rma_unit_domain(self, package_id=None, search=None, domain=None, match_type='all'):
        partner = request.env.user.partner_id
        base_domain = [('owner_id', '=', partner.commercial_partner_id.id)]
        
        if package_id:
            base_domain = expression.AND([base_domain, [('package_id', '=', int(package_id))]])
            
        if search:
            search_domain = [
                '|', '|', '|', '|', '|', '|',
                ('name', 'ilike', search),
                ('product_id.name', 'ilike', search),
                ('serial', 'ilike', search),
                ('imei', 'ilike', search),
                ('account_sku', 'ilike', search),
                ('account_ean13', 'ilike', search),
                ('account_fnsku', 'ilike', search),
                ('account_asin', 'ilike', search)
            ]
            base_domain = expression.AND([base_domain, search_domain])
            
        if domain:
            # Domain processing similar to portal_reception
            if match_type == 'any':
                advanced_domain = ['|'] * (len(domain) - 1) + domain
            else:
                advanced_domain = domain
            base_domain = expression.AND([base_domain, advanced_domain])
            
        return base_domain

    def _get_rma_unit_list_columns(self):
        return [
            {'id': 'name', 'label': _('RMA Number'), 'sortable': True},
            {'id': 'product', 'label': _('Product'), 'sortable': True},
            {'id': 'account_sku', 'label': _('Account SKU'), 'sortable': True, 'optional': 'show'},
            {'id': 'account_ean13', 'label': _('EAN13'), 'sortable': True, 'optional': 'hide'},
            {'id': 'account_fnsku', 'label': _('FNSKU'), 'sortable': True, 'optional': 'hide'},
            {'id': 'account_asin', 'label': _('ASIN'), 'sortable': True, 'optional': 'hide'},
            {'id': 'serial', 'label': _('Serial'), 'sortable': True, 'optional': 'show'},
            {'id': 'imei', 'label': _('IMEI'), 'sortable': True, 'optional': 'show'},
            {'id': 'location', 'label': _('Location'), 'sortable': True, 'optional': 'show'},
            {'id': 'condition', 'label': _('Condition'), 'sortable': True},
            {'id': 'state', 'label': _('State'), 'sortable': True},
            {'id': 'actions', 'label': _('Actions'), 'sortable': False, 'right': True},
        ]

    def _get_rma_unit_filters(self):
        return [
            {'id': 'all', 'label': _('All Units'), 'icon': 'fas fa-list', 'active': True, 'domain': []},
            {'id': 'condition_a', 'label': _('A. New'), 'icon': 'fas fa-check-circle', 'domain': [('condition', '=', 'A')]},
            {'id': 'condition_b', 'label': _('B. Used'), 'icon': 'fas fa-tools', 'domain': [('condition', '=', 'B')]},
            {'id': 'condition_c', 'label': _('C. Repair'), 'icon': 'fas fa-search', 'domain': [('condition', '=', 'C')]},
            {'id': 'condition_d', 'label': _('D. Discard'), 'icon': 'fas fa-trash-alt', 'domain': [('condition', '=', 'D')]},
            {'id': 'condition_e', 'label': _('E. On Hold'), 'icon': 'fas fa-pause', 'domain': [('condition', '=', 'E')]},
            {'id': 'to_inspect', 'label': _('To Inspect'), 'icon': 'fas fa-search', 'domain': [('state', '=', 'received')]},
            {'id': 'repairing', 'label': _('Repairing'), 'icon': 'fas fa-tools', 'domain': [('state', '=', 'repairing')]},
        ]

    def _get_rma_unit_advanced_search_fields(self):
        return [
            {'id': 'name', 'label': _('RMA Number'), 'type': 'char'},
            {'id': 'product_id', 'label': _('Product'), 'type': 'many2one', 'model': 'product.product'},
            {'id': 'serial', 'label': _('Serial'), 'type': 'char'},
            {'id': 'imei', 'label': _('IMEI'), 'type': 'char'},
            {'id': 'location_id', 'label': _('Location'), 'type': 'many2one', 'model': 'stock.location'},
            {'id': 'state', 'label': _('State'), 'type': 'selection', 'options': [
                ('draft', _('Draft')), ('received', _('Received')), ('inspected', _('Inspected')),
                ('repairing', _('Repairing')), ('done', _('Done')), ('scrapped', _('Scrapped'))
            ]},
            {'id': 'account_sku', 'label': _('Account SKU'), 'type': 'char'},
            {'id': 'account_ean13', 'label': _('EAN13'), 'type': 'char'},
            {'id': 'account_fnsku', 'label': _('FNSKU'), 'type': 'char'},
            {'id': 'account_asin', 'label': _('ASIN'), 'type': 'char'},
            {'id': 'create_date', 'label': _('Created On'), 'type': 'date'},
        ]

    @http.route('/account/rma/units', type='http', auth="user", website=True)
    def account_rma_units_action(self, package_id=None, **post):
        self._ensure_user_lang_context()
        values = self._get_admin_layout_values()
        
        limit = 20
        page = 1
        offset = (page - 1) * limit
        
        domain = self._get_rma_unit_domain(package_id)
        RmaUnit = request.env['rma.unit'].sudo()
        units = RmaUnit.search(domain, limit=limit, offset=offset, order='id desc')
        items_total = RmaUnit.search_count(domain)
        items_count = len(units)
        
        pagination_data = self._get_pagination_data(page, items_total, limit)
        
        list_columns = self._get_rma_unit_list_columns()
        list_filters = self._get_rma_unit_filters()

        batch_actions = [
            {'name': 'delete', 'label': _('Cancel'), 'icon': 'fas fa-trash-alt', 'color': 'bg-red-600 hover:bg-red-700'},
            {'name': 'export', 'label': _('Export'), 'icon': 'fas fa-file-export', 'color': 'bg-cyan-600 hover:bg-cyan-700'},
        ]

        values.update({
            'page_name': 'rma_units',
            'page_title': _('RMA Units'),
            'page_url': '/account/rma/units',
            'package_id': package_id,
            'list_columns': list_columns,
            'list_filters': list_filters,
            'units': units,
            'items_label': _('units'),
            'items_total': items_total or 0,
            'items_count': items_count or 0,
            **pagination_data,
            'tools_actions': [
                {'name': 'import', 'label': _('Import Units'), 'icon': 'fas fa-file-import', 'modal_id': 'dashboard-page-import-rma-modal'},
            ],
            'batch_actions': batch_actions,
            'advanced_search': json.dumps(self._get_rma_unit_advanced_search_fields()),
            'flatpickr': True,
            'select2': True,
        })
        
        return request.render("portal_rma.portal_rma_units_page_main", values)

    @http.route('/account/rma/units/reload', type='json', auth='user')
    def account_rma_units_reload(self, page=1, search='', domain=None, match_type='all', sort=None, order='asc', package_id=None, quick_filter=None, **kw):
        self._ensure_user_lang_context()
        limit = 20
        offset = (page - 1) * limit

        base_domain = self._get_rma_unit_domain(package_id, search, domain, match_type)
        
        # Apply filter domain if quick_filter is provided (similar to portal_reception)
        # Apply filter domain if quick_filter is provided (similar to portal_reception)
        if quick_filter:
            active_filter = next((f for f in self._get_rma_unit_filters() if f['id'] == quick_filter), None)
            if active_filter and active_filter.get('domain'):
                base_domain = expression.AND([base_domain, active_filter['domain']])

        order_by = 'id desc'
        if sort and sort in self.RMA_UNIT_FIELDS_MAPPING:
            order_by = f"{self.RMA_UNIT_FIELDS_MAPPING[sort]} {order}"

        RmaUnit = request.env['rma.unit'].sudo()
        units = RmaUnit.search(base_domain, limit=limit, offset=offset, order=order_by)
        items_total = RmaUnit.search_count(base_domain)
        items_count = len(units)

        pagination_data = self._get_pagination_data(page, items_total, limit)
        
        batch_actions = [
            {'name': 'delete', 'label': _('Cancel'), 'icon': 'fas fa-trash-alt', 'color': 'bg-red-600 hover:bg-red-700'},
            {'name': 'export', 'label': _('Export'), 'icon': 'fas fa-file-export', 'color': 'bg-cyan-600 hover:bg-cyan-700'},
        ]

        list_columns = self._get_rma_unit_list_columns()

        qweb = request.env['ir.qweb']
        return {
            'status': 'success',
            'list': qweb._render('portal_rma.portal_rma_list', {
                'units': units,
                'batch_actions': batch_actions,
                'list_columns': list_columns,
            }),
            'pager': qweb._render('portal_rma.portal_rma_pager', {
                'items_label': _('units'),
                'items_total': items_total or 0,
                'items_count': items_count or 0,
                **pagination_data
            }),
            'last_page': pagination_data.get('last_page', 0)
        }

    def _get_pagination_data(self, page, total, limit):
        last_page = (total + limit - 1) // limit if total > 0 else 0
        pages = []
        for p in range(max(1, page - 2), min(last_page, page + 2) + 1):
            pages.append({'page': p, 'active': p == page})
            
        return {
            'page': page,
            'total': total,
            'limit': limit,
            'last_page': last_page,
            'pages': pages,
            'items_count': total, # Compatibility
        }
    @http.route('/account/rma/units/download_template', type='http', auth='user')
    def account_rma_units_download_template(self, **kw):
        if not openpyxl:
            return request.not_found()
            
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "RMA Units Import"
        
        headers = ['Reference (Optional)', 'Product SKU or Name*', 'Serial Number', 'IMEI', 'Condition (A/B/C/D)', 'Notes']
        ws.append(headers)
        
        # Style headers
        from openpyxl.styles import Font, PatternFill
        header_row = next(ws.iter_rows(min_row=1, max_row=1))
        for cell in header_row:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

        fp = io.BytesIO()
        wb.save(fp)
        fp.seek(0)
        
        return request.make_response(fp.read(), headers=[
            ('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
            ('Content-Disposition', 'attachment; filename=rma_units_template.xlsx')
        ])

    @http.route('/account/rma/units/import', type='http', auth='user', methods=['POST'], csrf=False)
    def account_rma_units_import(self, file=None, package_id=None, **kw):
        if not openpyxl or not file:
            return request.make_json_response({'status': 'error', 'message': _('Import failed: Missing requirements.')})

        try:
            # Use local reference for type safety if needed, though simple check above is usually enough
            lib = openpyxl
            workbook = lib.load_workbook(file, data_only=True)
            sheet = workbook.active
            
            partner = request.env.user.partner_id.commercial_partner_id
            RmaUnit = request.env['rma.unit'].sudo()
            Product = request.env['product.product'].sudo()
            
            created_count = 0
            errors = []
            
            # Simple mapping for Condition
            cond_map = {'A': 'A', 'B': 'B', 'C': 'C', 'D': 'D'}

            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not any(row): continue
                
                # Slice to exactly 6 columns to avoid unpacking errors with extra empty columns
                row_data = row[:6]
                if len(row_data) < 6:
                    row_data = list(row_data) + [None] * (6 - len(row_data))
                
                ref, prod_ref, serial, imei, cond, notes = row_data
                
                if not prod_ref:
                    errors.append(_("Missing product reference at row %s") % str(row))
                    continue

                product = Product.search(['|', ('default_code', '=', str(prod_ref)), ('name', '=', str(prod_ref))], limit=1)
                if not product:
                    errors.append(_("Product %s not found") % str(prod_ref))
                    continue
                
                # Find mapping
                mapping = request.env['account.product.map'].sudo().search([
                    ('account_id.partner_id', '=', partner.id),
                    ('product_id', '=', product.id),
                    ('active', '=', True)
                ], limit=1)

                RmaUnit.create({
                    'product_id': product.id,
                    'owner_id': partner.id,
                    'account_product_map_id': mapping.id if mapping else False,
                    'serial': str(serial) if serial else False,
                    'imei': str(imei) if imei else False,
                    'condition': cond_map.get(str(cond).upper().strip(), 'E'),
                    'notes': str(notes) if notes else '',
                    'package_id': int(package_id) if package_id else False,
                    'state': 'received',
                    'received_date': fields.Datetime.now(),
                })
                created_count += 1

            return request.make_json_response({
                'status': 'success',
                'message': _('%d units imported successfully.') % created_count,
                'errors': errors
            })
        except Exception as e:
            _logger.error("RMA Import Error: %s", str(e))
            return request.make_json_response({'status': 'error', 'message': str(e)})

    # --- RMA Unit Creation & Mapping Endpoints ---

    @http.route('/account/rma/unit/create', type='json', auth='user')
    def account_rma_unit_create(self, **post):
        """Create new RMA units"""
        self._ensure_user_lang_context()
        partner_id = request.env.user.partner_id.commercial_partner_id.id
        
        products = json.loads(post.get('products') or '[]')
        if not products:
            return {'status': 'error', 'message': _('No products selected')}
            
        try:
            RmaUnit = request.env['rma.unit'].sudo()
            created_ids = []
            
            for p in products:
                qty = int(p.get('quantity', 1))
                
                # Find mapping
                mapping = request.env['account.product.map'].sudo().search([
                    ('account_id.partner_id', '=', partner_id),
                    ('product_id', '=', int(p['product_id'])),
                    ('active', '=', True)
                ], limit=1)

                for _ in range(qty):
                    unit = RmaUnit.create({
                        'product_id': int(p['product_id']),
                        'owner_id': partner_id,
                        'account_product_map_id': mapping.id if mapping else False,
                        'notes': post.get('notes'),
                        'state': 'received',
                        'received_date': fields.Datetime.now(),
                    })
                    created_ids.append(unit.id)
            
            return {
                'status': 'success', 
                'message': _('%d units created successfully') % len(created_ids),
                'ids': created_ids
            }
        except Exception as e:
            _logger.error("Error creating RMA units: %s", str(e))
            return {'status': 'error', 'message': str(e)}

    @http.route('/account/rma/product_maps', type='json', auth='user')
    def account_rma_product_maps(self, search='', **post):
        """Fetch existing product mappings for the client"""
        self._ensure_user_lang_context()
        partner_id = request.env.user.partner_id.commercial_partner_id
        
        AccountPartner = request.env['account.partner'].sudo()
        account_partner = AccountPartner.search([('partner_id', '=', partner_id.id)], limit=1)
        if not account_partner:
            return {'status': 'success', 'product_maps': []}

        domain = [('account_id', '=', account_partner.id), ('active', '=', True)]
        if search:
            domain = expression.AND([domain, [('name', 'ilike', search)]])
            
        maps = request.env['account.product.map'].sudo().search(domain, limit=50)
        
        result = []
        for m in maps:
            result.append({
                'id': m.id,
                'name': m.name,
                'account_sku': m.account_sku,
                'product_id': m.product_id.id
            })
        return {'status': 'success', 'product_maps': result}

    @http.route('/account/rma/unmapped_products', type='json', auth='user')
    def account_rma_unmapped_products(self, search='', **kw):
        """Fetch all products not yet mapped for the current client"""
        self._ensure_user_lang_context()
        partner_id = request.env.user.partner_id.commercial_partner_id
        
        AccountPartner = request.env['account.partner'].sudo()
        account_partner = AccountPartner.search([('partner_id', '=', partner_id.id)], limit=1)
        if not account_partner:
            return {'status': 'success', 'quants': []}

        ProductProduct = request.env['product.product'].sudo()
        mapped_product_ids = request.env['account.product.map'].sudo().search([
            ('account_id', '=', account_partner.id),
            ('active', '=', True)
        ]).product_id.ids

        domain = [('id', 'not in', mapped_product_ids)]
        if search:
            domain = expression.AND([
                domain,
                expression.OR([
                    [('name', 'ilike', search)],
                    [('default_code', 'ilike', search)]
                ])
            ])

        # Limit to 50 for performance
        products = ProductProduct.search(domain, limit=50)
        
        result = []
        for p in products:
            result.append({
                'id': p.id,
                'name': p.name,
                'code': p.default_code,
                # Link to the existing reception image endpoint since it's already there
                'image_url': f'/account/reception/product_image/{p.id}'
            })

        return {'status': 'success', 'quants': result}

    @http.route('/account/rma/product/map/quick_create', type='json', auth='user')
    def account_rma_product_map_quick_create(self, **post):
        """Quickly create an account.product.map for a product"""
        self._ensure_user_lang_context()
        partner_id = request.env.user.partner_id.commercial_partner_id
        product_id = int(post.get('product_id') or 0)
        if not product_id:
            return {'status': 'error', 'message': _('Product ID is required')}
        
        AccountPartner = request.env['account.partner'].sudo()
        account_partner = AccountPartner.search([('partner_id', '=', partner_id.id)], limit=1)
        if not account_partner:
            return {'status': 'error', 'message': _('Account partner not found')}

        try:
            ProductMap = request.env['account.product.map'].sudo()
            account_sku = post.get('account_sku') or ''
            
            # 1. Check if EXACT mapping exists (active or not)
            existing = ProductMap.with_context(active_test=False).search([
                ('account_id', '=', account_partner.id),
                ('account_sku', '=', account_sku)
            ], limit=1)

            if existing:
                if existing.product_id.id == int(product_id):
                    # Same product, just reactivate if needed
                    if not existing.active:
                        existing.write({'active': True})
                    mapping = existing
                else:
                    # SKU assigned to a DIFFERENT product
                    return {
                        'status': 'error', 
                        'message': _('SKU "%s" is already assigned to another product: %s') % (account_sku, existing.product_id.name)
                    }
            else:
                # 2. Create new mapping
                mapping = ProductMap.create({
                    'product_id': int(product_id),
                    'name': post.get('name'),
                    'account_sku': account_sku,
                    'account_id': account_partner.id,
                    'active': True
                })

            return {
                'status': 'success',
                'id': mapping.id,
                'name': mapping.name,
                'account_sku': mapping.account_sku,
                'product_name': mapping.product_id.name,
                'product_id': mapping.product_id.id
            }
        except Exception as e:
            return {'status': 'error', 'message': str(e)}

    @http.route('/account/rma/export', type='http', auth='user', methods=['GET', 'POST'])
    def account_rma_export(self, ids=None, **kw):
        """Export RMA units to Excel (XLSX) file."""
        self._ensure_user_lang_context()
        
        if not xlsxwriter:
            return request.make_response(
                _('Excel export not available. Please install xlsxwriter.'),
                headers=[('Content-Type', 'text/plain')]
            )
        
        RmaUnit = request.env['rma.unit'].sudo()
        partner_id = request.env.user.partner_id.commercial_partner_id
        
        domain = [('owner_id', '=', partner_id.id)]
        
        if ids:
            try:
                id_list = [int(i) for i in ids.split(',') if i.strip()]
                if id_list:
                    domain.append(('id', 'in', id_list))
            except ValueError:
                pass
        
        units = RmaUnit.search(domain, order='id desc')

        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {'in_memory': True})
        worksheet = workbook.add_worksheet(_('RMA Units'))
        
        header_format = workbook.add_format({
            'bold': True,
            'bg_color': '#0097a7',
            'font_color': 'white',
            'border': 1,
            'align': 'center',
            'valign': 'vcenter'
        })
        cell_format = workbook.add_format({
            'border': 1,
            'valign': 'vcenter'
        })
        date_format = workbook.add_format({
            'border': 1,
            'valign': 'vcenter',
            'num_format': 'dd/mm/yyyy hh:mm'
        })
        
        worksheet.set_column(0, 0, 15)  # RMA Number
        worksheet.set_column(1, 1, 30)  # Product
        worksheet.set_column(2, 2, 15)  # SKU
        worksheet.set_column(3, 3, 15)  # EAN13
        worksheet.set_column(4, 4, 15)  # Serial
        worksheet.set_column(5, 5, 10)  # State
        worksheet.set_column(6, 6, 20)  # Created On
        
        headers = [
            _('RMA Number'),
            _('Product'),
            _('Account SKU'),
            _('EAN13'),
            _('Serial'),
            _('State'),
            _('Created On'),
        ]
        
        for col, header in enumerate(headers):
            worksheet.write(0, col, header, header_format)
            
        row = 1
        for unit in units:
            worksheet.write(row, 0, unit.name, cell_format)
            worksheet.write(row, 1, unit.product_id.display_name, cell_format)
            worksheet.write(row, 2, unit.account_sku or '', cell_format)
            worksheet.write(row, 3, unit.account_ean13 or '', cell_format)
            worksheet.write(row, 4, unit.serial or '', cell_format)
            worksheet.write(row, 5, unit.state, cell_format)
            worksheet.write(row, 6, unit.create_date, date_format)
            row += 1
            
        workbook.close()
        output.seek(0)
        
        return request.make_response(output.read(), [
            ('Content-Type', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),
            ('Content-Disposition', 'attachment; filename=rma_units_export.xlsx;')
        ])
